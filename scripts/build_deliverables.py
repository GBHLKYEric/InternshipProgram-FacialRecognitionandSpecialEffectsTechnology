"""Generate professional tutorial PDFs, slides, and a complete code/use-site compendium.

Run with the project research/document environment. Reads source, never edits it.
"""
from __future__ import annotations
import ast
import hashlib
import html
import json
import re
import textwrap
from pathlib import Path

import markdown
from markdown.extensions.toc import slugify_unicode
from bs4 import BeautifulSoup, NavigableString
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Preformatted, Table, TableStyle, PageBreak
from reportlab.platypus import Image as PDFImage

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
REPORTS = ROOT / 'reports'
OUTPUT = ROOT.parent

CSS = '''body{font:17px/1.85 "Segoe UI","Microsoft YaHei",sans-serif;color:#1b2925;background:#f7f8f4;max-width:1040px;margin:auto;padding:44px 28px}h1{font-size:36px;line-height:1.4;color:#1c4939}h2{font-size:25px;margin-top:2.5em;border-bottom:1px solid #cdd7cf;padding-bottom:12px}h3{font-size:20px;margin-top:1.8em}a{color:#1f6550}p{margin:1em 0}code{font:14px/1.65 Consolas,"Microsoft YaHei",monospace;background:#eaf0ea;padding:2px 4px}pre{background:#142c25;color:#e4efdf;padding:22px;overflow:auto;border-radius:8px;line-height:1.7}pre code{background:transparent;color:inherit;white-space:pre}table{border-collapse:collapse;width:100%;font-size:14px}th,td{padding:12px;border:1px solid #d4ded6;vertical-align:top;overflow-wrap:anywhere}th{background:#e7eee8;text-align:left}.toc{font-size:14px;columns:2;column-gap:36px}.toc ul{padding-left:18px}blockquote{margin-left:0;padding-left:18px;border-left:3px solid #789c85;color:#40594b}img{max-width:100%}.meta{font-size:13px;color:#63756a}@media(max-width:600px){.toc{columns:1}body{padding:20px 16px}h1{font-size:28px}table{display:block;overflow:auto}}@media print{body{background:white;max-width:none;font-size:10pt;padding:0}.toc{columns:2}h2{break-before:page}pre{white-space:pre-wrap;overflow-wrap:anywhere;background:#f3f5f2;color:#111;font-size:8pt}pre code{white-space:pre-wrap}a{color:inherit}table{font-size:8pt}tr{break-inside:avoid}}'''


def rendered(source: str) -> str:
    source = re.sub(r'```mermaid\n.*?```', '\n**处理路径：** 图像或视频帧 → 检测与关键点 → 对齐 → 特征向量 → 余弦比较。检测与关键点也用于定位贴纸、美颜与三维重建。\n', source, flags=re.S)
    return markdown.markdown(source, extensions=['tables', 'fenced_code', 'toc', 'sane_lists'], extension_configs={'toc': {'slugify': slugify_unicode}})


def make_html(source: Path, destination: Path, title: str):
    body = rendered('[TOC]\n\n'+source.read_text(encoding='utf-8'))
    body = re.sub(r'href="([a-z-]+)\.md(?=[#"])', r'href="\1.html', body)
    body = body.replace('href="tutorial-zh.html', 'href="tutorial.html')
    destination.write_text(f'<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>{html.escape(title)}</title><style>{CSS}</style><body>{body}</body></html>', encoding='utf-8')


def source_files():
    result = []
    for folder in ['research','vision3d','scripts','tests','web','notebooks']:
        result.extend(p for p in (ROOT/folder).rglob('*') if p.is_file() and (p.suffix in {'.py','.html','.ipynb','.sh','.ps1','.yml','.yaml','.Dockerfile'} or p.name=='Dockerfile' or ('licenses' in p.parts and p.suffix=='.txt')) and '__pycache__' not in p.parts)
    result.extend(p for p in ROOT.iterdir() if p.is_file() and (p.suffix in {'.py','.txt','.ps1','.cmd','.toml','.yml'} or p.name in {'Dockerfile','.dockerignore','.gitignore','.gitattributes'}))
    result.append(ROOT/'models/registry.json')
    result.append(ROOT/'hello.py/Untitled-1.md')
    effective_config=REPORTS/'wider-mmdet-effective-config.py'
    if effective_config.exists(): result.append(effective_config)
    return sorted(set(result))


def use_of(path: Path) -> str:
    rel=path.relative_to(ROOT).as_posix()
    specific = {
        'desktop.py': '最终原生桌面前端；start.ps1默认启动。Tkinter窗口直接经OpenCV读取摄像头，调用app.Vision；原生三维窗口调用vision3d.reconstruct。',
        '启动桌面程序.cmd': 'Windows双击入口，调用同目录start.ps1并在失败时保留控制台。',
        'scripts/verify_desktop.py': 'desktop.py --self-test调用；实际Tk控件、摄像头采集/停止/重开/退出释放，以及本地三维和可选GAN检查。',
        'scripts/copy_source_delivery.py': '复制Git可见源文件到桌面目录，逐文件比对SHA256并生成交付清单；不删除目标现有文件。',
        'scripts/analyze_webcam.py': '识别网页/原生摄像头报告结构，独立复算帧数、FPS、分位数并绘制图表；原生报告额外分层说明检出人脸与未检出人脸的耗时。',
        'scripts/analyze_celeba.py': '读取真实202599张CelebA属性计数，生成全部40标签的频数图与5个目标属性摘要；由教程数据分析部分引用。',
        'scripts/record_desktop_demo.py': '只用公开NASA样图操作真实Tk控件与原生GL，录制桌面功能视频；5FPS编码不是算法速度。',
        'vision3d/native_viewer.py': 'desktop.py重建按钮调用start_viewer；独立spawn进程创建原生OpenGL窗口，完整网格、旋转与显式截图；--self-test验证三视角与退出。',
        'vision3d/licenses/pyglet-2.1.16-LICENSE.txt': '原生OpenGL窗口依赖pyglet的原始BSD许可证，随源码交付；不由本项目MIT许可覆盖。',
        'research/wider_training_report.py': '读取MMEngine真实逐步日志，输出CSV和损失曲线；区分进行中快照与最终完成。',
        'research/wider_watch.py': '监护已运行的完整WIDER训练，全部预测存在后执行官方协议与独立参考评估，不重启训练。',
        'research/wider_awake.py': '在本轮监护运行期间发出Windows防闲置休眠请求，完成、失败或超时后释放；不改电源计划且尊重用户手动睡眠。',
        'research/wider_checkpoint.py': '在对应MMDetection容器内读取本机最终checkpoint元数据，核对6440次更新及保存配置，不加载未知来源权重。',
        'research/wider_audit.py': '完整训练与两份评估完成后汇总配置、环境、checkpoint哈希及计时；要求临时电源请求已经释放。',
        'reports/wider-mmdet-effective-config.py': '本次真实运行保存的MMDetection完整有效配置；与checkpoint内嵌配置逐字一致。它保留旧版末轮重复验证行为，与后续修正入口分开记录。',
        'scripts/mmdet_thread_benchmark.py': '容器内用相同checkpoint和真实批次测试4/6/8线程训练步；微测不包含整条数据管线。',
        'scripts/mmdet_thread_benchmark_host.py': '临时暂停既有训练容器后执行线程微测，finally恢复同一个容器，保存调度证据。',
        'scripts/setup_anaconda.ps1': '下载校验官方Anaconda；显式接受条款参数后安装、创建独立conda环境并调用verify_anaconda.py。',
        'scripts/verify_anaconda.py': '检查实际conda解释器、PyTorch反向传播、YuNet推理、Jupyter内核执行和依赖一致性。',
        'scripts/probe_acceleration.py': '只读探测Intel GPU、OpenCL和ORT后端；不以设备名称冒充模型加速。',
        'scripts/setup_xpu.ps1': '建立独立Intel XPU环境并安装官方PyTorch XPU构建；不改主venv或驱动。',
        'scripts/verify_xpu.py': '实际XPU张量、NMS、GAN二阶反向传播和同步计时检查；真实训练前运行。',
        'research/ecosystem_simulation.py': 'parameter-server子命令模拟BytePS的两工作进程加权梯度聚合；inference-engine子命令以ORT图优化模拟ByteNN推理优化概念。',
        'research/wider_download.py': '完整WIDER数据下载、断点续传、SHA256与ZIP安全校验、完整COCO标注生成。',
        'research/wider_eval.py': '全部3226张WIDER验证图推理，以及官方难度mask/IoU/AP协议的Python计算；支持给定自训预测目录。',
        'research/wider_reference.py': '对同一WIDER预测调用固定哈希的OpenCV Zoo参考评估器，交叉检查AP。',
        'scripts/mmdet_full.py': '独立兼容Docker内的完整WIDER RetinaNet迁移训练、有限性检查、恢复、COCO验证和官方格式预测导出。',
        'research/stargan_deploy.py': '将实际StarGAN checkpoint导出到runs/stargan-deploy/generator.onnx，比较ORT数值；原生桌面和辅助网页共同使用该模型。',
    }
    if rel in specific: return specific[rel]
    if rel=='app.py': return '核心Vision检测/五点/验证/特效/属性编辑，被desktop.py直接调用；独立启动时也提供辅助网页HTTP服务。PDF任务2.4、4.3、7.2、9.2、9.3及系统集成。'
    if rel.startswith('web/'): return '由 app.py 的 GET / 返回；浏览器处理图像输入、效果控制、摄像头帧和双图验证。'
    if rel.startswith('vision3d/'): return '由 python -m vision3d.reconstruct 及其导出脚本调用；PDF任务8.2、8.3。'
    if rel.startswith('research/'): return '研究训练与评估命令使用；详见 research/README.md 和 docs/code-map.md。对应PDF任务3.2–7.3。'
    if rel.startswith('tests/'): return '回归检查；在项目根目录运行 python -m unittest discover -s tests。'
    if rel.startswith('notebooks/'): return 'Jupyter中的循序实验；PDF任务1.4、2.1、2.4和OpenCV入门。'
    if rel.startswith('scripts/'): return '模型下载、环境检查、数据准备或交付物生成；本文件开头docstring和下方函数索引提供具体入口。'
    if rel=='Dockerfile': return 'docker build --target hello 或 --target lab；PDF任务1.3。'
    if rel=='hello_world.py': return 'Docker hello镜像启动命令；PDF任务1.2/1.3。保留了旧仓库hello.py目录。'
    if rel=='hello.py/Untitled-1.md': return '用户原仓库保留的最初Hello World文件；当前Docker实际运行根目录hello_world.py，不执行此Markdown文件。'
    if rel=='start.ps1': return 'Windows环境准备及原生桌面启动入口；加-Web才打开辅助网页，-Port只用于网页模式。'
    if rel=='models/registry.json': return 'scripts/fetch_models.py 和 vision3d/reconstruct.py 读取；固定第三方模型来源、版本与校验值。'
    return '依赖固定、打包或版本控制配置；由 pip、Docker 或 Git 读取。'


def make_compendium():
    files=source_files()
    chunks=['# 全部项目源码与使用位置\n', '此文档逐字收录本项目源文件，并提供使用位置、符号行号、原文件链接与SHA256。它由 `scripts/build_deliverables.py` 从实际文件自动生成，避免手工粘贴漏代码。第三方Python库与预训练权重不复制进本文；它们的版本、官方来源与许可证见环境锁定文件、`models/registry.json`、`models/licenses/`及`docs/sources.md`。\n', '## 如何阅读\n先查用途，再打开对应文件。文中的代码是完整源码，不是删节片段。Notebook收录全部代码单元，图像输出在原始ipynb中。算法的专业解释见主教程。\n', '## 文件目录\n\n| 文件 | 源码行数 | 使用位置 |\n|---|---:|---|']
    total=0
    for p in files:
        raw=p.read_text(encoding='utf-8')
        count=len(raw.splitlines()); total+=count
        rel=p.relative_to(ROOT).as_posix()
        chunks.append(f'| `{rel}` | {count} | {use_of(p)} |')
    chunks.append(f'\n共 {len(files)} 个源文件，{total} 行文本（Notebook按JSON行统计）。\n')
    for p in files:
        rel=p.relative_to(ROOT).as_posix(); raw=p.read_text(encoding='utf-8')
        digest=hashlib.sha256(p.read_bytes()).hexdigest()
        chunks.extend([f'\n## {rel}\n',f'**使用位置：** {use_of(p)}\n',f'[GitHub中的文件](https://github.com/GBHLKYEric/InternshipProgram-FacialRecognitionandSpecialEffectsTechnology/blob/main/{rel})\n',f'SHA256：`{digest}`\n'])
        if p.suffix=='.py':
            try:
                tree=ast.parse(raw)
                symbols=[]
                def locate(parent, prefix=''):
                    for node in ast.iter_child_nodes(parent):
                        if isinstance(node,(ast.FunctionDef,ast.AsyncFunctionDef,ast.ClassDef)):
                            qualified=prefix+node.name
                            symbols.append(f'`{qualified}` 第{node.lineno}行')
                            locate(node,qualified+'.')
                        else:
                            locate(node,prefix)
                locate(tree)
                if symbols: chunks.append('**代码定位：** '+'；'.join(symbols)+'。\n')
            except SyntaxError: pass
        if p.suffix=='.ipynb':
            cells=json.loads(raw)['cells']
            for index,cell in enumerate(cells,1):
                if cell['cell_type']=='code':
                    code=''.join(cell['source']) if isinstance(cell['source'],list) else cell['source']
                    chunks.append(f'### 第 {index} 个单元\n\n```python\n{code}\n```\n')
        else:
            language={'.py':'python','.html':'html','.ps1':'powershell','.sh':'bash','.yml':'yaml','.toml':'toml'}.get(p.suffix,'text')
            # Four backticks safely contain any triple fences embedded in a builder string.
            chunks.append(f'````{language}\n{raw}\n````\n')
    path=DOCS/'code-compendium.md';path.write_text('\n'.join(chunks),encoding='utf-8')
    make_html(path,DOCS/'code-compendium.html','全部源码与使用位置')
    (REPORTS/'code-inventory.json').write_text(json.dumps({'files':len(files),'lines':total,'paths':[p.relative_to(ROOT).as_posix() for p in files]},indent=2),encoding='utf-8')


def pdf_from_markdown(source: Path, destination: Path):
    font=Path('C:/Windows/Fonts/msyh.ttc')
    if not font.exists(): raise FileNotFoundError('PDF generation requires Microsoft YaHei, or adjust font path')
    if 'Chinese' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('Chinese',str(font),subfontIndex=0))
        pdfmetrics.registerFontFamily('Chinese',normal='Chinese',bold='Chinese',italic='Chinese',boldItalic='Chinese')
    styles=getSampleStyleSheet()
    for style in styles.byName.values(): style.fontName='Chinese'
    body=ParagraphStyle('BodyCJK',fontName='Chinese',fontSize=10.2,leading=17,spaceAfter=8,wordWrap='CJK',allowOrphans=0,allowWidows=0)
    code=ParagraphStyle('CodeCJK',fontName='Chinese',fontSize=8.2,leading=12,spaceAfter=10,backColor=colors.HexColor('#eef3ee'),borderPadding=7)
    heading=ParagraphStyle('HeadingCJK',parent=body,fontSize=16,leading=23,spaceBefore=20,spaceAfter=12,keepWithNext=True,textColor=colors.HexColor('#214b39'))
    small=ParagraphStyle('SmallCJK',parent=body,fontSize=8,leading=12)
    soup=BeautifulSoup(rendered(source.read_text(encoding='utf-8')),'html.parser')
    story=[]
    def inline(node):
        def child_markup(child):
            if isinstance(child, NavigableString):
                return html.escape(str(child)).replace('\n', '<br/>')
            contents = ''.join(child_markup(part) for part in child.children)
            if child.name in {'strong', 'b'}:
                return f'<b>{contents}</b>'
            if child.name in {'em', 'i'}:
                return f'<i>{contents}</i>'
            if child.name == 'a' and child.get('href', '').startswith(('https://', 'http://')):
                return f'<link href="{html.escape(child["href"], quote=True)}" color="#28674c">{contents}</link>'
            return contents
        return ''.join(child_markup(child) for child in node.children)
    for node in soup.children:
        if not getattr(node,'name',None): continue
        name=node.name
        if name in {'h1','h2','h3','h4'}:
            if name=='h2' and node.get_text().startswith('第') and story: story.append(PageBreak())
            style=heading if name!='h1' else ParagraphStyle('TitleCJK',parent=heading,fontSize=24,leading=34,spaceAfter=20)
            following=node.find_next_sibling()
            if following is not None and following.name=='table':
                # Let a long table split on this page instead of keeping the
                # heading together with the entire multi-page table.
                style=ParagraphStyle('HeadingBeforeTable',parent=style,keepWithNext=False)
            story.append(Paragraph(inline(node),style))
        elif name=='pre':
            lines=[]
            for line in node.get_text().splitlines():
                lines.extend(textwrap.wrap(line,width=84,expand_tabs=False,replace_whitespace=False,drop_whitespace=False) or [''])
            story.append(Preformatted('\n'.join(lines),code))
        elif name=='table':
            rows=[[Paragraph(inline(cell),small) for cell in row.find_all(['th','td'])] for row in node.find_all('tr')]
            if rows:
                table=Table(rows,colWidths=[(174*mm)/len(rows[0])]*len(rows[0]),repeatRows=1,hAlign='LEFT')
                table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e5eee7')),('GRID',(0,0),(-1,-1),.4,colors.HexColor('#c8d7cc')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
                story.extend([table,Spacer(1,10)])
        elif name in {'ul','ol'}:
            for i,item in enumerate(node.find_all('li',recursive=False),1): story.append(Paragraph((f'{i}. ' if name=='ol' else '• ')+inline(item),body))
        elif name=='p' and node.find('img'):
            for picture in node.find_all('img'):
                relative = picture.get('src', '')
                candidate = (source.parent/relative).resolve()
                if relative and ROOT in candidate.parents and candidate.is_file():
                    from PIL import Image
                    with Image.open(candidate) as im:
                        width, height = im.size
                    scale = min(174*mm/width, 130*mm/height)
                    story.extend([PDFImage(str(candidate), width=width*scale, height=height*scale), Spacer(1, 8)])
                if picture.get('alt'): story.append(Paragraph(html.escape(picture['alt']), small))
        elif name=='hr': story.append(Spacer(1,12))
        else:
            paragraph = Paragraph(inline(node),body)
            following = node.find_next_sibling()
            if following is not None and following.name == 'pre':
                paragraph.keepWithNext = True
            story.append(paragraph)
    def footer(canvas,doc):
        canvas.setFont('Chinese',8);canvas.setFillColor(colors.HexColor('#62776a'))
        canvas.drawString(18*mm,12*mm,'Face Vision Lab  ·  可复现实验与专业教程')
        canvas.drawRightString(192*mm,12*mm,str(doc.page))
    document=SimpleDocTemplate(str(destination),pagesize=(210*mm,297*mm),leftMargin=18*mm,rightMargin=18*mm,topMargin=18*mm,bottomMargin=22*mm,title=source.stem,author='Face Vision Lab')
    document.build(story,onFirstPage=footer,onLaterPages=footer)


def slides():
    metrics=json.loads((REPORTS/'demo-metrics.json').read_text(encoding='utf-8'))
    lfw_path=REPORTS/'lfw-sface-largest.json'
    lfw=json.loads(lfw_path.read_text(encoding='utf-8')) if lfw_path.exists() else json.loads((REPORTS/'lfw-sface.json').read_text(encoding='utf-8'))
    opt=json.loads((REPORTS/'arcface-pilot/comparison.json').read_text(encoding='utf-8'))
    wider=json.loads((REPORTS/'wider-mmdet-full-official.json').read_text(encoding='utf-8'))
    assert wider['full_validation'] and wider['counts']['images']==3226
    wider_ap=' / '.join(f"{wider['metrics'][name]['ap']:.5f}" for name in ['easy','medium','hard'])
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    entries=[
      ('人脸视觉 从原理到可运行系统',['本地 CPU 应用 · 训练与评估管线 · 3D 重建','面向高三学生的可复现学习项目','独立教育项目，与字节跳动无隶属关系'],None),
      ('范围与验收口径',['PDF 含环境、检测、关键点、识别、优化、GAN、3D 与动态特效','预训练推理、代码冒烟测试、真实数据实验是三种不同证据','98.5% 是目标，必须由完整协议的实测报告支持'],None),
      ('系统如何处理一张图像',['图像 → YuNet 检测框与五点 → 对齐 → SFace 特征 → 余弦比较','关键点 → 几何变换 → 眼镜、皇冠、美颜、美妆','最终为原生 Tk 桌面窗口；OpenCV 直接读取相机，图像留在内存'],None),
      ('真实原生桌面界面',['直接启动本机窗口，六种显示模式；图片与摄像头输入','后台线程读取/推理，主线程更新 Tk 窗口','停止、重开和退出释放摄像头已实测'], 'reports/desktop-window.png'),
      ('环境与可重复性',['Windows 11 · Core Ultra 5 225H · 32 GiB RAM · Intel Arc 130T','Anaconda 2026.07-1 / conda 26.5.3；项目 Python 3.12.14','主 venv、conda、XPU 和 MMDetection Docker 分开验证'],None),
      ('人脸识别与 ArcFace',['对齐后的人脸 → ResNet50 → L2 归一化特征','ArcFace 训练正类 logit：s cos(θ + m)，使类间更易区分','P-K 采样与 batch-hard triplet 辅助；实测反向传播及保存重读'],None),
      ('真实训练：结果尚未达到目标',['52 个身份 / 133 张训练图；3 epochs；测试身份名无交集','平均训练损失：40.112 → 36.925 → 36.126',f'自训 FP32 完整 LFW：{opt["accuracy"]["fp32"]["accuracy_mean"]*100:.4f}%；小数据短训未收敛'],None),
      ('LFW 的真实验证结果',[f'完整 6,000 对 / 10 折；准确率 {lfw.get("accuracy_mean",0)*100:.2f}%','阈值仅由其余训练折确定；测试折不选阈值','预训练 SFace 基线；不能归属于本项目从零训练的 ResNet50'],None),
      ('一次失败怎样变成有效实验',['首轮严格单脸策略：72.25%；全部失败都来自多脸图','原图包含背景人脸；应用验证与数据集主体选择规则不同','保留失败报告，按最大主体规则独立重跑；不删除困难样本'], 'reports/lfw-detection-audit.jpg'),
      ('动态特效与性能',[f'演示视频 18 秒 / 360 帧；{metrics["frames_with_face"]} 帧检测到人脸',f'混合特效管线中位数 {metrics["pipeline_ms_median"]:.2f} ms，P95 {metrics["pipeline_ms_p95"]:.2f} ms','公开静态样本经仿射运动；不是摄像头实测或移动端性能'], 'reports/effect-all.jpg'),
      ('本机摄像头的独立实测',['原生桌面：640×480，组合特效；20.0601 秒 / 149 帧','提交帧率 7.4277 FPS；延迟中位 136.49 ms / P95 159.36 ms','电脑同时训练模型；完整口径见报告，不称流畅度或手机达标'],None),
      ('第二次摄像头测量怎样读',['316 帧 / 20.0267 秒，15.779 FPS；只有 49 帧检出人脸','有脸中位 161.46 ms；无脸中位 40.60 ms，统计组成不同','不能把整体 FPS 变高解释成优化；保留两轮原始报告'],None),
      ('从单张照片重建三维',['3DDFA_V2：预测 62 参数，重建 38,365 顶点 / 76,073 三角面','原生 OpenGL 4.6 全网格；0°、+35°、−35°实测','单目统计估计，无真实尺度；没有 3D 真值误差测试'], 'reports/desktop-3d.png'),
      ('真实 CelebA 与 StarGAN 微调',['完整下载 202,599 张图像包；按清单提取本次所需样本','从作者预训练模型继续：最终保留 1,200 D / 240 G 更新','有限规模微调；不是从零全量训练或论文级复现'],None),
      ('生成质量评估的正确边界',['固定 768 源图与 768 参考图，图片不重叠不等于身份不重叠','匹配 XPU 构建：FID 41.2728 → 40.9804；IS 3.0881 → 3.0980','小样本一次比较；没有显著性或身份保持结论'],None),
      ('把真实生成模型接入桌面',['固定输入：128×128 RGB 与 5 个目标标签','PyTorch / ONNX Runtime 最大像素数值差 1.97e−6','按钮暂停摄像头，展示输入裁剪与合成结果；不是实时 GAN'], 'reports/desktop-attributes.png'),
      ('真实训练模型的部署对比',[f'批量 8 / 2 线程：FP32 {opt["fp32"]["median_ms"]:.2f} ms；int8 {opt["dynamic_linear_int8"]["median_ms"]:.2f} ms',f'ONNX {opt["onnx"]["median_ms"]:.2f} ms；该设置下两者均未加速',f'int8 完整 LFW：{opt["accuracy"]["dynamic_linear_int8"]["accuracy_mean"]*100:.4f}%；微小变化不能证明更准确'],None),
      ('研究训练的完整路径',['WIDER → COCO 标注 → MMDetection；300-W → 关键点 → NME','身份文件夹 → ArcFace；CelebA 属性 → StarGAN → FID/IS','受数据授权、版本依赖与算力约束；逐项状态见验收矩阵'],None),
      ('WIDER 完整验证与交叉检查',['全部 3,226 张验证图；31 张零检测保留为空','YuNet AP：easy 0.88442 / medium 0.86568 / hard 0.75040','与 OpenCV Zoo 参考评估三项差为 0；预训练基线不归属自训模型'],None),
      ('完整 WIDER 迁移训练的实测结果',['12,880 张训练图；6,440 次更新；全部 3,226 张验证图',f'easy / medium / hard AP：{wider_ap}','COCO 预训练迁移、冻结主干、最大边 320；不代表论文收敛结果'],'reports/wider-mmdet-full-official-pr.png'),
      ('训练曲线与失败记录怎样读',['学习率 0.0005，250 步 warmup，梯度范数裁剪 10','曲线是实际日志窗口均值；下降不等于验证成绩达标','首次 NaN、待机时间、参考评估交叉检查均保留证据'],'reports/wider-mmdet-training-curves.png'),
      ('Intel Arc 的实际 XPU 路径',['独立 PyTorch 2.14 XPU 环境；不改主 CPU 环境或驱动','张量、NMS、GAN 二阶梯度及优化器更新实际通过','完整生成器微测中位 23.66 ms；形状/负载有限定，不外推全训练'],None),
      ('BytePS 与 ByteNN 概念实验',['两个真实工作进程按 3/8 与 5/8 聚合梯度；与参考误差 4.47e−8','ORT 图优化：batch1 本次中位 34.37 → 25.27 ms','两项都标为模拟；没有冒充内部 SDK 或移动端部署'],None),
      ('火山引擎官方样例体验',['官方人像融合页内置方案交互：两张示例输入与融合输出','未上传用户照片；截图只用于记录官方样例体验','展示可能是预计算，不能据此宣称完成鉴权云 API 部署'], 'reports/volcengine-experience.png'),
      ('版本与环境问题怎样解决',['Anaconda 长路径 → 物理短目录，不改全局 PATH 和注册表','旧检测栈 → 独立 Docker；新 CPU/XPU → 各自固定构建','StarGAN 旧运行统计造成偏色 → 对齐作者逐图 InstanceNorm'],None),
      ('专业学习路线与源码导航',['18 课：Python → 图像/向量 → 网络 → 评估 → 生成/3D → 部署','每课：准确概念、直观解释、运行步骤、练习与答案','全部源码汇编附用途、调用入口、函数行号和哈希'],None),
      ('开源交付与尚未完成的验收',['公开仓库在 README 中提供；完整源码另存桌面并核对 SHA256','MIT 适用于本项目代码；权重与数据分别遵守原许可','自训 98.5% 目标、受限数据和真实手机验收仍逐项保留'],None),
    ]
    for number,(title,bullets,picture) in enumerate(entries,1):
        slide=prs.slides.add_slide(prs.slide_layouts[6]);bg=slide.background.fill;bg.solid();bg.fore_color.rgb=RGBColor.from_string('F6F7F1')
        def textbox(x,y,w,h,text,size,color='20382D',bold=False):
            shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=shape.text_frame;tf.word_wrap=True
            p=tf.paragraphs[0];p.text=text;p.font.name='Microsoft YaHei';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
            return shape
        textbox(.65,.35,11.7,.35,'FACE VISION LAB  /  PROJECT REVIEW',11,'647A69')
        textbox(.65,.95,12,1,title,30,bold=True)
        if picture and (ROOT/picture).exists():
            from PIL import Image
            with Image.open(ROOT/picture) as im: ratio=im.width/im.height
            boxw,boxh=6.2,4.7;w=min(boxw,boxh*ratio);h=w/ratio
            slide.shapes.add_picture(str(ROOT/picture),Inches(6.5+(boxw-w)/2),Inches(2.05+(boxh-h)/2),width=Inches(w),height=Inches(h))
            width=5.2
        else: width=11.6
        for i,bullet in enumerate(bullets):
            textbox(.7,2.25+i*1.25,width,1.15,bullet,22 if picture else 25)
        textbox(.65,7,10,.25,'可复现的代码、明确的证据、可理解的原理',10,'687E70')
        textbox(12,7,.6,.25,f'{number:02d}',10,'687E70')
    path=OUTPUT/'项目汇报.pptx';prs.save(path)
    outline=['# 项目汇报逐页大纲\n', '此大纲由实际PPT内容同步生成；数字的完整协议与限制见项目报告。\n']
    for number, (title, bullets, picture) in enumerate(entries, 1):
        outline.append(f'## {number}. {title}\n')
        outline.extend(f'- {bullet}' for bullet in bullets)
        if picture: outline.append(f'\n配图：`{picture}`。')
        outline.append('\n')
    (DOCS/'presentation-outline.md').write_text('\n'.join(outline), encoding='utf-8')
    make_html(DOCS/'presentation-outline.md', DOCS/'presentation-outline.html', '项目汇报逐页大纲')
    check=Presentation(path)
    assert len(check.slides)==len(entries)
    for slide in check.slides:
        for shape in slide.shapes:
            assert shape.left>=0 and shape.top>=0 and shape.left+shape.width<=check.slide_width+10000 and shape.top+shape.height<=check.slide_height+10000
    print(f'Created {len(entries)} slides')


def main():
    make_compendium()
    for source in DOCS.glob('*.md'):
        if source.name=='code-compendium.md': continue
        dest=DOCS/('tutorial.html' if source.name=='tutorial-zh.md' else source.stem+'.html')
        make_html(source,dest,source.stem)
    pdf_from_markdown(DOCS/'tutorial-zh.md',OUTPUT/'高三零基础专业教程.pdf')
    for stem,name in [('project-report','项目总结报告.pdf'),('issues-and-fixes','执行问题与解决记录.pdf'),('continued-experiments','继续实验专业教程.pdf'),('desktop-guide','本地桌面程序使用与代码教程.pdf')]:
        if (DOCS/(stem+'.md')).exists(): pdf_from_markdown(DOCS/(stem+'.md'),OUTPUT/name)
    slides()
    print('Generated tutorial, reports, slides and complete code compendium.')


if __name__=='__main__': main()
